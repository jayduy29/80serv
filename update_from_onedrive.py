# -*- coding: utf-8 -*-
"""
Script cập nhật dữ liệu dây điện cho ứng dụng 80SERV trực tiếp từ link chia sẻ Microsoft OneDrive (.xlsx).
Cách dùng:
    python update_from_onedrive.py "https://1drv.ms/x/c/..."
hoặc chỉ cần chạy:
    python update_from_onedrive.py
"""

import sys
if hasattr(sys.stdout, 'reconfigure'):
    try:
        sys.stdout.reconfigure(encoding='utf-8')
        sys.stderr.reconfigure(encoding='utf-8')
    except Exception:
        pass

import os
import json
import urllib.request
import http.cookiejar

def download_onedrive_file(share_url, output_path="checklist_finall.xlsx"):
    share_url = share_url.strip()
    if not share_url:
        raise ValueError("Đường link OneDrive không được để trống.")
    
    if "download=1" not in share_url:
        delim = "&" if "?" in share_url else "?"
        download_url = share_url + delim + "download=1"
    else:
        download_url = share_url

    print(f"[*] Đang kết nối tới OneDrive: {share_url[:70]}...")

    cj = http.cookiejar.CookieJar()
    opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(cj))
    opener.addheaders = [
        ('User-Agent', 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36')
    ]

    try:
        resp = opener.open(download_url, timeout=30)
        data = resp.read()
    except Exception as e:
        raise RuntimeError(f"Lỗi khi tải file từ OneDrive: {e}")

    if not data.startswith(b'PK\x03\x04'):
        raise RuntimeError("Dữ liệu tải về không phải file Excel (.xlsx) hợp lệ. Vui lòng kiểm tra quyền chia sẻ của link OneDrive (chọn 'Bất kỳ ai có liên kết đều có thể xem').")

    with open(output_path, "wb") as f:
        f.write(data)

    print(f"[+] Tải file thành công: {output_path} ({len(data):,} bytes)")
    return output_path

def clean_status(val):
    if val is None:
        return ""
    if isinstance(val, float) and val.is_integer():
        return str(int(val))
    s = str(val).strip()
    if s.endswith(".0"):
        s = s[:-2]
    return s

def parse_xlsx_cables(xlsx_path):
    print(f"[*] Đang phân tích file Excel: {xlsx_path}...")
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    
    sheets_list = []
    counts = {}
    total_obj = {'name': 'TỔNG', 'gid': '0'}
    all_cables = {}

    for i, name in enumerate(wb.sheetnames):
        ws = wb[name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        h_idx = -1
        for r_i, r in enumerate(rows[:5]):
            first_val = str(r[0] or '').strip().lower()
            if 'cable' in first_val or 'mã cáp' in first_val or 'số dây' in first_val:
                h_idx = r_i
                break
        if h_idx == -1:
            h_idx = 1 if len(rows) > 1 and rows[1][0] and 'cable' in str(rows[1][0]).lower() else 0
        data_rows = rows[h_idx+1:]
        
        cables = []
        for r in data_rows:
            if not r or r[0] is None:
                continue
            c_num = str(r[0]).strip()
            if not c_num or c_num.lower().startswith(('cable', 'tổng', 'số dây', 'stt')):
                continue
            source = str(r[1]).strip() if len(r) > 1 and r[1] is not None else ""
            target = str(r[3]).strip() if len(r) > 3 and r[3] is not None else ""
            c_type = str(r[4]).strip() if len(r) > 4 and r[4] is not None else ""
            spec = str(r[5]).strip() if len(r) > 5 and r[5] is not None else ""
            keo = clean_status(r[6] if len(r) > 6 else None)
            dau = clean_status(r[7] if len(r) > 7 else None)
            contractor = str(r[9]).strip() if len(r) > 9 and r[9] is not None else ""
            cables.append([c_num, source, target, c_type, spec, contractor, keo, dau])

        all_cables[name] = cables
        if name.lower() in ('total', 'tổng', 'tong'):
            total_obj = {'name': name, 'gid': '0'}
        else:
            sheets_list.append({'name': name, 'gid': str(i)})
            counts[name] = len(cables)

    print(f"[+] Đã đọc: {len(sheets_list)} sheet tủ điện, tổng số dây: {len(all_cables.get(total_obj['name'], [])):,}")
    return total_obj, sheets_list, counts, all_cables

def update_index_html(total_obj, sheets_list, counts, all_cables, html_path="index.html"):
    print(f"[*] Đang cập nhật vào {html_path}...")
    with open(html_path, "r", encoding="utf-8") as f:
        content = f.read()

    new_seed = {
        'total': total_obj,
        'sheets': sheets_list,
        'counts': counts
    }
    seed_json = json.dumps(new_seed, ensure_ascii=False, separators=(',', ':'))
    cables_json = json.dumps(all_cables, ensure_ascii=False, separators=(',', ':'))

    seed_idx = content.find('const WIRE_SEED=')
    if seed_idx == -1:
        raise RuntimeError("Không tìm thấy const WIRE_SEED trong index.html")
    after_seed = content.find('let WIRE_SHEETS=', seed_idx)
    if after_seed == -1:
        raise RuntimeError("Không tìm thấy let WIRE_SHEETS= sau WIRE_SEED trong index.html")

    replacement = f"const WIRE_SEED={seed_json};\nconst WIRE_SEED_CABLES={cables_json};\n\n"
    content = content[:seed_idx] + replacement + content[after_seed:]

    with open(html_path, "w", encoding="utf-8", newline="") as f:
        f.write(content)

    print(f"[✔] HOÀN TẤT! Đã cập nhật thành công toàn bộ dữ liệu vào {html_path}!")

DEFAULT_ONEDRIVE_URL = "https://1drv.ms/x/c/e898f0bace036960/IQD2rR5GS2FjTYh2peM3EgROAY2D9geLBA0U7pVrTl7Nqmg?e=NuDBrO"

if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1].strip():
        onedrive_url = sys.argv[1].strip()
    else:
        onedrive_url = DEFAULT_ONEDRIVE_URL
        print(f"[*] Dùng link OneDrive mặc định: {DEFAULT_ONEDRIVE_URL}")

    if not onedrive_url:
        print("Lỗi: Chưa nhập link.")
        sys.exit(1)

    try:
        xlsx_file = download_onedrive_file(onedrive_url)
        total_obj, sheets_list, counts, all_cables = parse_xlsx_cables(xlsx_file)
        update_index_html(total_obj, sheets_list, counts, all_cables)
    except Exception as err:
        print(f"\n[X] LỖI: {err}")
        sys.exit(1)
